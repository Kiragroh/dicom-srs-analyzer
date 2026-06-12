"""
Export brain and normal-brain VxGy metrics to Excel.

Example:
    python brain_normal_export.py --data-root ../data --output output/brain_normal_metrics.xlsx

The exporter uses RTSTRUCT + RTDOSE directly. CT images are not required.
"""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
import pydicom
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from brain_normal_metrics import (
    active_numbers_from_plan_label,
    global_brain_dose_metrics,
    local_brain_dose_metrics,
    mask_volume_cc,
    numbers_after_prefix,
    rounded_d99_prescription_gy,
    select_brain_structure,
    structure_mask_on_dose_grid,
)
from dicom_io import load_dose, load_plan_meta, load_structures


DEFAULT_LOCAL_THRESHOLDS = "5,10,12"
DEFAULT_GLOBAL_THRESHOLDS = "5,8,10,12"
DEFAULT_LOCAL_MARGIN_MM = 15.0


@dataclass
class DicomPlanSet:
    case_id: str
    rtstruct_path: Path
    rtplan_path: Path
    rtdose_path: Path


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def case_hash(*parts: object) -> str:
    text = "|".join(clean_text(part) for part in parts)
    return hashlib.sha1(text.encode("utf-8", errors="ignore")).hexdigest()[:12]


def parse_thresholds(text: str) -> list[float]:
    values = []
    for item in text.split(","):
        item = item.strip()
        if not item:
            continue
        value = float(item)
        if value <= 0:
            raise ValueError("Dose thresholds must be positive.")
        values.append(value)
    if not values:
        raise ValueError("At least one dose threshold is required.")
    return values


def referenced_plan_uid(ds) -> str:
    seq = getattr(ds, "ReferencedRTPlanSequence", None)
    if seq:
        return clean_text(getattr(seq[0], "ReferencedSOPInstanceUID", ""))
    return ""


def referenced_structure_set_uid(ds) -> str:
    seq = getattr(ds, "ReferencedStructureSetSequence", None)
    if seq:
        return clean_text(getattr(seq[0], "ReferencedSOPInstanceUID", ""))
    return ""


def discover_plan_sets(root: Path) -> list[DicomPlanSet]:
    records = []
    for path in root.rglob("*.dcm"):
        try:
            ds = pydicom.dcmread(str(path), stop_before_pixels=True, force=True)
        except Exception:
            continue
        modality = clean_text(getattr(ds, "Modality", ""))
        if modality not in {"RTSTRUCT", "RTPLAN", "RTDOSE"}:
            continue
        top_folder = path.relative_to(root).parts[0] if path != root else path.parent.name
        records.append(
            {
                "group": top_folder,
                "path": path,
                "modality": modality,
                "sop_uid": clean_text(getattr(ds, "SOPInstanceUID", "")),
                "ref_plan_uid": referenced_plan_uid(ds),
                "ref_rs_uid": referenced_structure_set_uid(ds),
            }
        )

    if not records:
        return []

    plan_sets: list[DicomPlanSet] = []
    for _group, group in pd.DataFrame(records).groupby("group", dropna=False):
        rs_rows = group[group["modality"].eq("RTSTRUCT")]
        rp_rows = group[group["modality"].eq("RTPLAN")]
        rd_rows = group[group["modality"].eq("RTDOSE")]
        if rs_rows.empty or rp_rows.empty or rd_rows.empty:
            continue

        for _, rd_row in rd_rows.iterrows():
            matching_rp = rp_rows
            if rd_row["ref_plan_uid"]:
                by_uid = rp_rows[rp_rows["sop_uid"].eq(rd_row["ref_plan_uid"])]
                if not by_uid.empty:
                    matching_rp = by_uid
            rp_row = matching_rp.iloc[0]

            matching_rs = rs_rows
            if rp_row["ref_rs_uid"]:
                by_uid = rs_rows[rs_rows["sop_uid"].eq(rp_row["ref_rs_uid"])]
                if not by_uid.empty:
                    matching_rs = by_uid

            plan_sets.append(
                DicomPlanSet(
                    case_id=f"case_{len(plan_sets) + 1:03d}",
                    rtstruct_path=Path(matching_rs.iloc[0]["path"]),
                    rtplan_path=Path(rp_row["path"]),
                    rtdose_path=Path(rd_row["path"]),
                )
            )
    return plan_sets


def plan_target_prescription_gy(rtplan_path: Path) -> float:
    ds = pydicom.dcmread(str(rtplan_path), stop_before_pixels=True, force=True)
    for item in getattr(ds, "DoseReferenceSequence", []):
        if clean_text(getattr(item, "DoseReferenceType", "")).upper() != "TARGET":
            continue
        try:
            value = float(getattr(item, "TargetPrescriptionDose", ""))
        except Exception:
            continue
        if value > 0:
            return value
    return float("nan")


def plan_fraction_count(rtplan_path: Path) -> float:
    ds = pydicom.dcmread(str(rtplan_path), stop_before_pixels=True, force=True)
    for group in getattr(ds, "FractionGroupSequence", []):
        try:
            value = float(getattr(group, "NumberOfFractionsPlanned", ""))
        except Exception:
            continue
        if value > 0:
            return value
    return float("nan")


def build_target_maps(structures) -> tuple[dict[int, object], dict[int, object], pd.DataFrame]:
    ptv_by_number = {}
    gtv_by_number = {}
    rows = []
    for structure in structures:
        ptv_numbers = numbers_after_prefix(structure.name, "PTV")
        gtv_numbers = numbers_after_prefix(structure.name, "GTV")
        composite = len(ptv_numbers) > 1 or len(gtv_numbers) > 1
        if len(ptv_numbers) == 1 and not composite:
            ptv_by_number[ptv_numbers[0]] = structure
        if len(gtv_numbers) == 1 and not composite:
            gtv_by_number[gtv_numbers[0]] = structure
        rows.append(
            {
                "StructureName": structure.name,
                "DICOMType": structure.dicom_type,
                "ROINumber": structure.roi_number,
                "ContourCount": len(structure.contours),
                "PTVNumber": ptv_numbers[0] if len(ptv_numbers) == 1 else np.nan,
                "GTVNumber": gtv_numbers[0] if len(gtv_numbers) == 1 else np.nan,
                "CompositeTarget": composite,
            }
        )
    return ptv_by_number, gtv_by_number, pd.DataFrame(rows)


def run_plan_set(
    plan_set: DicomPlanSet,
    *,
    local_thresholds_gy: list[float],
    global_thresholds_gy: list[float],
    local_margin_mm: float,
    allow_external_fallback: bool,
) -> dict[str, pd.DataFrame]:
    structures = load_structures(str(plan_set.rtstruct_path))
    dose = load_dose(str(plan_set.rtdose_path))
    meta = load_plan_meta(str(plan_set.rtplan_path))
    if dose is None or meta is None:
        raise RuntimeError(f"Could not load RTDOSE/RTPLAN for {plan_set.case_id}")

    brain = select_brain_structure(structures, allow_external_fallback=allow_external_fallback)
    if brain is None:
        raise RuntimeError(f"No brain contour found for {plan_set.case_id}")

    plan_label = clean_text(meta.plan_label)
    active_numbers = active_numbers_from_plan_label(plan_label)
    rx_plan = plan_target_prescription_gy(plan_set.rtplan_path)
    fractions = plan_fraction_count(plan_set.rtplan_path)

    ptv_by_number, gtv_by_number, inventory = build_target_maps(structures)
    paired_numbers = sorted(set(ptv_by_number).intersection(gtv_by_number))

    mask_cache: dict[str, np.ndarray] = {}

    def mask_for(structure):
        key = structure.name
        if key not in mask_cache:
            mask_cache[key] = structure_mask_on_dose_grid(structure, dose)
        return mask_cache[key]

    brain_mask = mask_for(brain.structure)
    ptv_union = np.zeros_like(brain_mask)
    gtv_union = np.zeros_like(brain_mask)
    lesion_rows = []

    for number in paired_numbers:
        ptv = ptv_by_number[number]
        gtv = gtv_by_number[number]
        ptv_mask = mask_for(ptv)
        gtv_mask = mask_for(gtv)
        rx_d99 = rounded_d99_prescription_gy(ptv_mask, dose)
        active_by_label = number in active_numbers if active_numbers else False
        active_by_d99 = bool(np.isfinite(rx_d99) and rx_d99 >= 10)
        active = active_by_label or (not active_numbers and active_by_d99)
        if active:
            ptv_union |= ptv_mask
            gtv_union |= gtv_mask

        row = {
            "CaseId": plan_set.case_id,
            "PlanHash": case_hash(plan_set.rtplan_path.name, plan_label),
            "PlanLabel": plan_label,
            "TargetNumber": number,
            "PTV": ptv.name,
            "GTV": gtv.name,
            "Fractions": fractions,
            "RxPlanGy": rx_plan,
            "RxD99RoundedGy": rx_d99,
            "PTVVolumeCc": mask_volume_cc(ptv_mask, dose),
            "GTVVolumeCc": mask_volume_cc(gtv_mask, dose),
            "ActiveByPlanLabel": active_by_label,
            "ActiveByD99": active_by_d99,
            "ActiveForPlan": active,
            "BrainStructure": brain.structure.name,
            "BrainSource": brain.source,
            "BrainExternalFallback": brain.is_external_fallback,
        }
        row.update(
            local_brain_dose_metrics(
                dose=dose,
                brain_mask=brain_mask,
                ptv_mask=ptv_mask,
                gtv_mask=gtv_mask,
                ptv_structure=ptv,
                thresholds_gy=local_thresholds_gy,
                margin_mm=local_margin_mm,
            )
        )
        lesion_rows.append(row)

    plan_row = {
        "CaseId": plan_set.case_id,
        "PlanHash": case_hash(plan_set.rtplan_path.name, plan_label),
        "PlanLabel": plan_label,
        "Fractions": fractions,
        "RxPlanGy": rx_plan,
        "ActiveTargetNumbers": ",".join(str(n) for n in paired_numbers if (n in active_numbers or not active_numbers)),
        "BrainStructure": brain.structure.name,
        "BrainSource": brain.source,
        "BrainExternalFallback": brain.is_external_fallback,
    }
    plan_row.update(
        global_brain_dose_metrics(
            dose=dose,
            brain_mask=brain_mask,
            ptv_union_mask=ptv_union,
            gtv_union_mask=gtv_union,
            thresholds_gy=global_thresholds_gy,
        )
    )

    inventory.insert(0, "CaseId", plan_set.case_id)
    inventory["UsedAsBrain"] = inventory["StructureName"].eq(brain.structure.name)

    qa = pd.DataFrame(
        [
            {
                "CaseId": plan_set.case_id,
                "PlanLabel": plan_label,
                "RTSTRUCT": plan_set.rtstruct_path.name,
                "RTPLAN": plan_set.rtplan_path.name,
                "RTDOSE": plan_set.rtdose_path.name,
                "PTVGTVPairs": len(paired_numbers),
                "ActiveTargetsFromPlanLabel": ",".join(str(n) for n in sorted(active_numbers)),
                "BrainStructure": brain.structure.name,
                "BrainSource": brain.source,
                "BrainExternalFallback": brain.is_external_fallback,
                "DoseGridShape": "x".join(str(v) for v in dose.dose_grid.shape),
                "DoseSpacingMm": ",".join(f"{float(v):g}" for v in dose.spacing),
                "CTUsed": False,
            }
        ]
    )

    return {
        "lesion_brain_normal": pd.DataFrame(lesion_rows),
        "plan_brain_normal": pd.DataFrame([plan_row]),
        "structure_inventory": inventory,
        "qa": qa,
    }


def concat_results(results: list[dict[str, pd.DataFrame]]) -> dict[str, pd.DataFrame]:
    sheets = {}
    for name in ["lesion_brain_normal", "plan_brain_normal", "structure_inventory", "qa"]:
        frames = [r[name] for r in results if name in r and not r[name].empty]
        sheets[name] = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    sheets["metrics_long"] = make_long_metrics(sheets["lesion_brain_normal"], sheets["plan_brain_normal"])
    sheets["README"] = pd.DataFrame(
        [
            {"Item": "Purpose", "Value": "Export VxGy metrics for brain, brain minus PTV, and brain minus GTV."},
            {"Item": "Geometry", "Value": "RTSTRUCT contours are voxelized directly on the RTDOSE grid; CT images are not used."},
            {"Item": "Brain selection", "Value": "Whole Brain / Brain / Hirn are preferred; EXTERNAL/BODY can be used as fallback."},
            {"Item": "Normal brain", "Value": "BrainMinusPTV subtracts the PTV mask; BrainMinusGTV subtracts the GTV mask."},
            {"Item": "Local metrics", "Value": "Local metrics are calculated in a radial expansion around each PTV bounding box."},
            {"Item": "Global metrics", "Value": "Global metrics use the full selected brain mask and active PTV/GTV union masks."},
            {"Item": "Prescription helper", "Value": "RxD99RoundedGy is included for traceability but VxGy thresholds are absolute Gy thresholds."},
        ]
    )
    return sheets


def make_long_metrics(lesions: pd.DataFrame, plans: pd.DataFrame) -> pd.DataFrame:
    rows = []
    if not lesions.empty:
        ids = ["CaseId", "PlanLabel", "TargetNumber", "PTV", "GTV"]
        for _, row in lesions.iterrows():
            for col in lesions.columns:
                if col in ids:
                    continue
                rows.append({**{key: row.get(key, "") for key in ids}, "Scope": "lesion", "Metric": col, "Value": row.get(col)})
    if not plans.empty:
        ids = ["CaseId", "PlanLabel"]
        for _, row in plans.iterrows():
            for col in plans.columns:
                if col in ids:
                    continue
                rows.append({**{key: row.get(key, "") for key in ids}, "Scope": "plan", "Metric": col, "Value": row.get(col)})
    return pd.DataFrame(rows)


def write_workbook(sheets: dict[str, pd.DataFrame], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    order = ["README", "lesion_brain_normal", "plan_brain_normal", "metrics_long", "structure_inventory", "qa"]
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet in order:
            sheets.get(sheet, pd.DataFrame()).to_excel(writer, sheet_name=sheet[:31], index=False)

    wb = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, col in enumerate(ws.columns, start=1):
            values = [clean_text(c.value) for c in list(col)[:200]]
            width = min(max([len(v) for v in values] + [10]) + 2, 50)
            ws.column_dimensions[get_column_letter(idx)].width = width
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export brain / normal-brain VxGy metrics from SRS DICOM data.")
    parser.add_argument("--data-root", type=Path, default=Path("data"), help="Folder containing DICOM patient/plan folders.")
    parser.add_argument("--output", type=Path, default=Path("output") / "brain_normal_metrics.xlsx")
    parser.add_argument("--local-thresholds", default=DEFAULT_LOCAL_THRESHOLDS, help="Comma-separated local VxGy thresholds in Gy.")
    parser.add_argument("--global-thresholds", default=DEFAULT_GLOBAL_THRESHOLDS, help="Comma-separated global VxGy thresholds in Gy.")
    parser.add_argument("--local-margin-mm", type=float, default=DEFAULT_LOCAL_MARGIN_MM)
    parser.add_argument("--no-external-fallback", action="store_true", help="Fail when no brain contour exists instead of using EXTERNAL/BODY.")
    args = parser.parse_args()

    plan_sets = discover_plan_sets(args.data_root)
    if not plan_sets:
        raise RuntimeError(f"No RTSTRUCT/RTPLAN/RTDOSE sets found under {args.data_root}")

    local_thresholds = parse_thresholds(args.local_thresholds)
    global_thresholds = parse_thresholds(args.global_thresholds)
    results = [
        run_plan_set(
            plan_set,
            local_thresholds_gy=local_thresholds,
            global_thresholds_gy=global_thresholds,
            local_margin_mm=args.local_margin_mm,
            allow_external_fallback=not args.no_external_fallback,
        )
        for plan_set in plan_sets
    ]
    sheets = concat_results(results)
    write_workbook(sheets, args.output)

    print(f"plans={len(plan_sets)}")
    print(f"lesion_rows={len(sheets['lesion_brain_normal'])}")
    print(f"output={args.output}")


if __name__ == "__main__":
    main()

